from __future__ import annotations

from datetime import date
from pathlib import Path

from lxml import etree

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import CONTENT_TYPE as CT


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "outputs" / "manual"
OUTPUT_FILE = OUTPUT_DIR / "黄金对冲分析网站_风险热力图与操作手册_20260814.docx"


# compact_reference_guide preset + editorial_cover template.
TOKENS = {
    "page_width_in": 8.5,
    "page_height_in": 11.0,
    "margin_in": 1.0,
    "content_width_dxa": 9360,
    "table_indent_dxa": 120,
    "body_font": "Arial Unicode MS",
    "mono_font": "Aptos Mono",
    "body_size_pt": 11,
    "body_color": "222222",
    "muted": "536273",
    "navy": "17365D",
    "blue": "2E74B5",
    "deep_blue": "1F4D78",
    "light_blue": "E8EEF5",
    "lightest_blue": "F5F8FC",
    "green": "1F7A4D",
    "light_green": "EDF7F0",
    "gold": "B7791F",
    "light_gold": "FFF4E5",
    "red": "9B1C1C",
    "light_red": "FDECEC",
    "border": "A8B6C7",
}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TOKENS["content_width_dxa"]))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TOKENS["table_indent_dxa"]))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            cant_split = OxmlElement("w:cantSplit")
            cant_split.set(qn("w:val"), "true")
            tr_pr.append(cant_split)
        for index, cell in enumerate(row.cells):
            width = widths_dxa[min(index, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_east_asia_font(run, font_name: str) -> None:
    run.font.name = font_name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font_name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font_name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font_name)


def set_field(run, instruction: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_hyperlink(paragraph, label: str, url: str) -> None:
    relation_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), TOKENS["blue"])
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    text = OxmlElement("w:t")
    text.text = label
    run.extend([r_pr, text])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def paragraph_border_and_fill(paragraph, fill: str, border: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "20")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), border)
    p_bdr.append(left)
    p_pr.append(p_bdr)


def create_numbering(doc: Document, ordered: bool) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "multilevel")
    abstract.append(multi)
    for level in range(3):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(level))
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "decimal" if ordered else "bullet")
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), f"%{level + 1}." if ordered else "•")
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), str(540 + level * 360))
        tabs.append(tab)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), str(540 + level * 360))
        ind.set(qn("w:hanging"), "270")
        p_pr.extend([tabs, ind])
        lvl.extend([start, num_fmt, lvl_text, suff, p_pr])
        abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num_id)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int, level: int = 0) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(TOKENS["page_width_in"])
    section.page_height = Inches(TOKENS["page_height_in"])
    section.top_margin = Inches(TOKENS["margin_in"])
    section.bottom_margin = Inches(TOKENS["margin_in"])
    section.left_margin = Inches(TOKENS["margin_in"])
    section.right_margin = Inches(TOKENS["margin_in"])
    section.header_distance = Inches(0.42)
    section.footer_distance = Inches(0.42)

    normal = doc.styles["Normal"]
    normal.font.name = TOKENS["body_font"]
    normal.font.size = Pt(TOKENS["body_size_pt"])
    normal.font.color.rgb = RGBColor.from_string(TOKENS["body_color"])
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), TOKENS["body_font"])
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.widow_control = True

    heading_specs = {
        "Title": (30, TOKENS["navy"], 0, 12),
        "Subtitle": (15, TOKENS["muted"], 0, 12),
        "Heading 1": (16, TOKENS["blue"], 18, 10),
        "Heading 2": (13, TOKENS["blue"], 14, 7),
        "Heading 3": (12, TOKENS["deep_blue"], 10, 5),
    }
    for style_name, (size, color, before, after) in heading_specs.items():
        style = doc.styles[style_name]
        style.font.name = TOKENS["body_font"]
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = style_name != "Subtitle"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), TOKENS["body_font"])
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = TOKENS["body_font"]
        style.font.size = Pt(TOKENS["body_size_pt"])
        style._element.rPr.rFonts.set(qn("w:eastAsia"), TOKENS["body_font"])
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    run = hp.add_run("黄金对冲分析网站  |  风险热力图与操作手册")
    set_east_asia_font(run, TOKENS["body_font"])
    run.font.size = Pt(8)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(TOKENS["muted"])
    hp.add_run("\t")
    run = hp.add_run("内部培训资料")
    set_east_asia_font(run, TOKENS["body_font"])
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(TOKENS["gold"])
    tabs = hp.paragraph_format.tab_stops
    tabs.add_tab_stop(Inches(6.5))

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_after = Pt(0)
    run = fp.add_run("DinoSignal Quant Risk  ·  内部使用  ·  第 ")
    set_east_asia_font(run, TOKENS["body_font"])
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(TOKENS["muted"])
    page_run = fp.add_run()
    set_field(page_run, "PAGE")
    run = fp.add_run(" 页")
    set_east_asia_font(run, TOKENS["body_font"])
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(TOKENS["muted"])

    doc.core_properties.title = "黄金对冲分析网站：风险热力图与操作手册"
    doc.core_properties.subject = "GLD 与 XAUT 期权仓位、风险热力图、操作流程和培训讲稿"
    doc.core_properties.author = "DinoSignal Quant Risk"
    doc.core_properties.keywords = "GLD, XAUT, options, risk heatmap, hedge, Greeks"

    # python-docx's default font table does not register custom CJK fonts.
    # LibreOffice will otherwise render Simplified Chinese as blanks even when
    # every run contains w:rFonts, so declare the selected family explicitly.
    font_part = next(
        (part for part in doc.part.package.parts if part.content_type == CT.WML_FONT_TABLE),
        None,
    )
    if font_part is not None:
        root = parse_xml(font_part.blob)
        if not any(node.get(qn("w:name")) == TOKENS["body_font"] for node in root.findall(qn("w:font"))):
            font = OxmlElement("w:font")
            font.set(qn("w:name"), TOKENS["body_font"])
            charset = OxmlElement("w:charset")
            charset.set(qn("w:val"), "86")
            family = OxmlElement("w:family")
            family.set(qn("w:val"), "swiss")
            pitch = OxmlElement("w:pitch")
            pitch.set(qn("w:val"), "variable")
            signature = OxmlElement("w:sig")
            signature.set(qn("w:usb0"), "E0002AFF")
            signature.set(qn("w:usb1"), "C0007843")
            signature.set(qn("w:usb2"), "00000009")
            signature.set(qn("w:usb3"), "00000000")
            signature.set(qn("w:csb0"), "00040001")
            signature.set(qn("w:csb1"), "00000000")
            font.extend([charset, family, pitch, signature])
            root.append(font)
            font_part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


class ManualBuilder:
    def __init__(self, doc: Document):
        self.doc = doc
        self.bullet_num = create_numbering(doc, ordered=False)
        self.order_num = create_numbering(doc, ordered=True)

    def p(self, text: str = "", bold_lead: str | None = None, align=None):
        paragraph = self.doc.add_paragraph()
        if align is not None:
            paragraph.alignment = align
        if bold_lead and text.startswith(bold_lead):
            lead = paragraph.add_run(bold_lead)
            lead.bold = True
            set_east_asia_font(lead, TOKENS["body_font"])
            tail = paragraph.add_run(text[len(bold_lead):])
            set_east_asia_font(tail, TOKENS["body_font"])
        else:
            run = paragraph.add_run(text)
            set_east_asia_font(run, TOKENS["body_font"])
        return paragraph

    def heading(self, text: str, level: int = 1):
        return self.doc.add_heading(text, level=level)

    def bullets(self, items: list[str], level: int = 0):
        for item in items:
            paragraph = self.doc.add_paragraph(style="List Bullet")
            apply_numbering(paragraph, self.bullet_num, level)
            run = paragraph.add_run(item)
            set_east_asia_font(run, TOKENS["body_font"])

    def steps(self, items: list[str], level: int = 0):
        order_num = create_numbering(self.doc, ordered=True)
        for item in items:
            paragraph = self.doc.add_paragraph(style="List Number")
            apply_numbering(paragraph, order_num, level)
            run = paragraph.add_run(item)
            set_east_asia_font(run, TOKENS["body_font"])

    def callout(self, title: str, body: str, kind: str = "info"):
        palette = {
            "info": (TOKENS["light_blue"], TOKENS["blue"]),
            "positive": (TOKENS["light_green"], TOKENS["green"]),
            "warning": (TOKENS["light_gold"], TOKENS["gold"]),
            "risk": (TOKENS["light_red"], TOKENS["red"]),
        }
        fill, border = palette[kind]
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.14)
        paragraph.paragraph_format.right_indent = Inches(0.04)
        paragraph.paragraph_format.space_before = Pt(5)
        paragraph.paragraph_format.space_after = Pt(7)
        paragraph.paragraph_format.keep_together = True
        paragraph_border_and_fill(paragraph, fill, border)
        lead = paragraph.add_run(f"{title}  ")
        lead.bold = True
        lead.font.color.rgb = RGBColor.from_string(border)
        set_east_asia_font(lead, TOKENS["body_font"])
        tail = paragraph.add_run(body)
        set_east_asia_font(tail, TOKENS["body_font"])
        return paragraph

    def formula(self, text: str, note: str | None = None):
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.22)
        paragraph.paragraph_format.right_indent = Inches(0.1)
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph_border_and_fill(paragraph, TOKENS["lightest_blue"], TOKENS["border"])
        run = paragraph.add_run(text)
        set_east_asia_font(run, TOKENS["mono_font"])
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor.from_string(TOKENS["deep_blue"])
        if note:
            note_run = paragraph.add_run(f"\n{note}")
            set_east_asia_font(note_run, TOKENS["body_font"])
            note_run.font.size = Pt(9)
            note_run.font.color.rgb = RGBColor.from_string(TOKENS["muted"])
        return paragraph

    def table(self, headers: list[str], rows: list[list[str]], widths_dxa: list[int]):
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        set_table_geometry(table, widths_dxa)
        header = table.rows[0]
        set_repeat_table_header(header)
        for index, label in enumerate(headers):
            cell = header.cells[index]
            set_cell_shading(cell, TOKENS["light_blue"])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(label)
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(TOKENS["navy"])
            set_east_asia_font(run, TOKENS["body_font"])
        for row_index, values in enumerate(rows):
            row = table.add_row()
            for index, value in enumerate(values):
                cell = row.cells[index]
                if row_index % 2:
                    set_cell_shading(cell, "F8FAFC")
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(str(value))
                set_east_asia_font(run, TOKENS["body_font"])
                run.font.size = Pt(9.5)
        # Apply fixed widths, cell margins, and no-split behavior after every
        # data row exists. Applying this only at table creation would affect
        # the header row and could leave a split row fragment on the next page.
        set_table_geometry(table, widths_dxa)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)
        return table

    def page_break(self):
        p = self.doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)


def build_manual() -> Document:
    doc = Document()
    configure_document(doc)
    b = ManualBuilder(doc)

    # Editorial cover.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Inches(0.68)
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run("INSTITUTIONAL OPTIONS RISK OPERATING GUIDE")
    set_east_asia_font(run, TOKENS["mono_font"])
    run.font.size = Pt(9)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(TOKENS["gold"])

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.keep_with_next = True
    run = title.add_run("黄金对冲分析网站")
    set_east_asia_font(run, TOKENS["body_font"])
    subtitle = doc.add_paragraph(style="Subtitle")
    run = subtitle.add_run("GLD + XAUT 仓位、风险热力图与对冲决策使用手册")
    set_east_asia_font(run, TOKENS["body_font"])

    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(6)
    rule.paragraph_format.space_after = Pt(18)
    paragraph_border_and_fill(rule, "FFFFFF", TOKENS["gold"])

    meta = b.p("版本：2026-08-14｜适用角色：交易、风险、运营、管理与培训｜系统：DinoSignal Gold Options Hedge")
    meta.paragraph_format.space_after = Pt(18)
    b.callout(
        "一句话定位",
        "把分散在到期日、行权价、标的与数据源中的期权风险，压缩为一个可执行的二维决策界面：先找风险，再判断数据，再决定对冲、减仓、展期或等待。",
        "positive",
    )
    b.p("本手册只把当前已经默认显示的能力写成“可直接使用”；隐藏功能、需要管理员开放的功能和未来路线图均单独标注。市场数据、模型和公式仅为决策支持，不替代交易确认、经纪商合约主数据、资金校验及双人复核。")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.keep_together = True
    lead = p.add_run("公网入口：")
    lead.bold = True
    set_east_asia_font(lead, TOKENS["body_font"])
    add_hyperlink(p, "风险热力图", "https://gold-options-hedge.onrender.com/matrix")
    p.add_run("  ·  ")
    add_hyperlink(p, "仓位管理", "https://gold-options-hedge.onrender.com/positions")
    p.add_run("  ·  ")
    add_hyperlink(p, "设置", "https://gold-options-hedge.onrender.com/settings")

    b.page_break()

    b.heading("阅读导航", 1)
    b.p("如果只有三分钟，先读第 1、3、4 节；如果要执行日常流程，读第 5、6、8 节；如果要向他人讲解，直接使用第 11 节讲稿；如果要规划下一版，读第 9、10 节。")
    b.table(
        ["章节", "回答的问题", "主要读者"],
        [
            ["1–2", "这个系统到底解决什么；边界在哪里", "全员"],
            ["3–4", "默认热力图看什么、怎么看、怎么操作", "交易 / 风险"],
            ["5–6", "仓位如何导入、数据如何更新、公式如何理解", "交易 / 运营 / 量化"],
            ["7–8", "盘前、盘中、到期日前如何形成标准动作", "交易 / 风控负责人"],
            ["9–10", "隐藏能力、当前问题与升级路线", "量化 / 技术 / 管理层"],
            ["11–12", "四类讲解稿与第一性原理总结", "培训 / 汇报"],
            ["附录", "术语、默认值、检查清单", "日常查阅"],
        ],
        [1150, 5100, 3110],
    )

    b.heading("1. 重点先看：它对交易与对冲有什么帮助", 1)
    b.heading("1.1 六个最重要的帮助", 2)
    b.bullets([
        "把复杂风险压缩成一张图。Expiry × Strike 是期权风险最自然的坐标；颜色编码当前 Metric，白色边框标出真实持仓，交易员无需逐行读表就能看到风险集中区。",
        "同时区分“单张风险”和“仓位风险”。Unit Delta 告诉你一张合约对价格的局部敏感度；Total Delta 把数量、合约乘数和黄金换算纳入，避免把 0.80 × 1 张误认为比 0.05 × 1,000 张更重要。",
        "把风险放回市场位置。Spot 数值精确到两位小数，并映射到最近 Strike；当 Spot 越界时仍显示 Above/Below Range，使 ITM/ATM/OTM 位置不会因坐标范围而丢失。",
        "把数据质量变成风险的一部分。MISSING、STALE、WARN、FAIL 不会静默变成 0；若所选 Metric 的计算链遇到缺失，方格保留可交易边框但不参与色标。",
        "让比较具有统一规则。默认绿—黄—红色标表达低—中—高；数量与名义本金只在真实持仓之间比较，避免完整期权链中的零仓位稀释色阶。",
        "形成可复核的操作链。Excel 导入或手工维护仓位后，更新市场数据、读热力图、Hover 核验、点击看详情、导出更新后的持仓，构成同一套可追溯流程。",
    ])
    b.callout(
        "交易员的三秒钟问题",
        "现在风险最大在哪里？这个颜色是仓位规模、Delta 还是波动率？Spot 离该 Strike 多远？这个数据够不够新、够不够完整？",
        "info",
    )

    b.heading("1.2 它不是什么", 2)
    b.bullets([
        "不是自动下单系统：热力图给出风险排序和上下文，不代表交易指令。",
        "不是经纪商账本：网站仓位必须与券商、交易所或托管对账，尤其是 adjusted contract、行权与交割。",
        "不是绝对实时保证：速度取决于数据授权和上游接口；页面时间戳与状态必须先于数值解读。",
        "不是全模型真值：Greeks 与情景重估依赖输入、定价模型和约定，极端行情下需结合流动性、跳跃和波动率曲面。",
    ])

    b.heading("2. 系统定位与当前默认信息架构", 1)
    b.p("当前默认导航只展示三个入口：风险热力图、仓位管理、设置。风险热力图和仓位管理用于日常交易；设置用于管理默认筛选项、标签、Hover、详情弹窗和页面入口。其他页面属于可选能力，默认隐藏或受管理员密码保护。")
    b.table(
        ["默认入口", "核心任务", "完成标准"],
        [
            ["风险热力图", "识别风险簇、Spot 位置、持仓/可交易合约、缺失数据", "能定位最重要格子并说清其 Metric、数据状态与下一步"],
            ["仓位管理", "录入、Excel 导入、刷新、核对与导出", "网站与来源仓位一致，更新后计算和导出同步"],
            ["设置", "决定热力图上能看到什么、页面入口和安全开关", "默认界面足够简洁，同时保留按角色开放能力"],
        ],
        [1500, 3900, 3960],
    )
    b.callout("默认展示原则", "默认值优先服务快速决策：GLD、Call、Notional Size、无数字标签、Risk Hover；更复杂的筛选、情景、卡片和错误榜单不占据首屏。", "positive")

    b.heading("3. 风险热力图：默认屏幕如何读", 1)
    b.heading("3.1 坐标与视觉语法", 2)
    b.bullets([
        "横轴是 Expiry，到期日由左向右展开；纵轴是 Strike，默认按低到高排列。每个格子代表一个 Expiry × Strike 聚合单元。",
        "背景色表达当前选择的 Metric：默认由低值绿色、中心黄色到高值红色。颜色只回答“相对大小”，具体数值由标签、Hover 或详情确认。",
        "白色较粗实线边框代表有真实持仓；青色细实线代表期权链中可交易但没有持仓的合约。边框不覆盖背景色判断。",
        "持仓格可显示 G/X（GLD/XAUT）、C/P（Call/Put）与 L/S/W/M/F（数据状态）角标；具体显示项由设置页控制。",
        "Spot 在最近 Strike 轴线上显示，数值固定为两位小数；Spot 超出当前 Strike 范围时，顶部显示 Above Range 或 Below Range，并可 Center Spot。",
        "如果一个可交易格子的所选 Metric 缺少必要输入，该格仍有四边实线，但背景留空且不参加色标，不能把空白误读为低风险。",
    ])

    b.heading("3.2 默认控件与正确用法", 2)
    b.table(
        ["控件", "默认", "它回答什么", "典型动作"],
        [
            ["Underlying", "GLD", "看哪个标的的风险面", "切到 XAUT 比较同一 Metric；固定色标后再横向比较"],
            ["C/P", "Call", "看 Call 或 Put", "分别看两侧风险，不把符号与结构混在一起"],
            ["Metric", "Notional Size", "颜色到底代表什么", "先看规模，再切 Total Delta、Unit Delta 与 IV"],
            ["Label", "None", "哪些格子直接显示数字", "Held Metric 只标持仓；Top/Bottom 15% 仅按需开启"],
            ["Hover", "Risk", "鼠标经过展示哪组字段", "快速核验 Unit Delta、Total Delta 与所选 Metric"],
            ["Range", "自动", "当前色标上下限", "跨 GLD/XAUT 比较时固定同一 Metric 的上下限"],
            ["Strike 排序", "升序", "纵轴方向", "按个人读图习惯反转，不改变数值"],
            ["方格尺寸", "紧凑", "细节与全局的取舍", "缩小看全貌，放大做逐格核查"],
            ["Fit All", "可见", "是否尽量一屏装下", "先全景找热点，再退出 Fit All 放大"],
            ["Full Screen", "可见", "是否扩大有效画布", "会议、盘中监控和 100+ 合约时使用"],
        ],
        [1180, 950, 3000, 4230],
    )

    b.heading("3.3 默认可选 Metric：不要混淆问题", 2)
    b.table(
        ["Metric", "业务含义", "色标范围", "最适合回答"],
        [
            ["Notional Size", "Qty × 标的现价 × 每张对应标的数量", "只用真实持仓", "哪一腿对应的标的名义本金最大"],
            ["Qty", "净持仓数量；热力图按绝对值比较", "只用真实持仓", "风险是否主要来自仓位数量"],
            ["Unit Delta", "每单位期权对标的价格变化的一阶敏感度", "可用合约", "哪一张合约最接近方向性敞口"],
            ["Total Delta", "Unit Delta × Qty × 合约乘数 × 黄金换算", "有值的合约", "真实仓位的黄金方向风险集中在哪里"],
            ["Mark IV", "Mark Price 对应的隐含波动率", "有有效 IV 的合约", "波动率曲面哪里贵/便宜或异常"],
            ["Bid IV / Ask IV", "买一/卖一对应的隐含波动率", "有有效 IV 的合约", "两侧报价隐波与可成交区间"],
            ["Ask-Bid IV Spread", "Ask IV − Bid IV", "有双边 IV 的合约", "哪一段波动率流动性最差"],
        ],
        [1450, 3300, 1500, 3110],
    )
    b.callout("本次新增", "Label 中新增 Bottom 15%。Top 15% 与 Bottom 15% 默认都隐藏，需先在设置页开放；排名按当前 Metric 的绝对值，并排除 MISSING。", "info")

    b.heading("4. 热力图标准使用法", 1)
    b.heading("4.1 开盘前：90 秒建立风险地图", 2)
    b.steps([
        "确认顶部市场数据更新时间和状态。先看 Source / As-of / STALE / MISSING，再看颜色；数据不可信时停止交易结论。",
        "保持默认 GLD + Call + Notional Size。找白色边框中的红色或黄色聚集区，确认名义规模集中在哪些 Expiry × Strike。",
        "切换到 Total Delta。识别仓位真实方向风险是否与名义规模一致；若不一致，通常来自不同 Delta、数量或合约乘数。",
        "切换 Unit Delta。区分“合约本身敏感”与“因为 Qty 大而总风险大”。",
        "切换 Mark IV、Bid IV、Ask IV 与 IV Spread。检查风险集中区是否同时流动性差或隐波异常。",
        "核对 Spot 两位小数及最近 Strike，判断风险簇处于 ITM、ATM 还是 OTM 附近。",
        "切到 Put，再按同样顺序检查；最后切 XAUT。需要跨标的比较时，为同一 Metric 使用相同的固定上下限。",
    ])

    b.heading("4.2 从颜色到动作：六步证据链", 2)
    b.steps([
        "定位：找到白框高色阶格，或可交易链上的异常 IV/Spread 格。",
        "识别：确认 Underlying、Call/Put、Expiry、Strike 与 Spot 位置。",
        "核数：Hover 查看所选 Metric、Unit Delta、Total Delta 和数据状态。",
        "下钻：点击格子查看完整详情，包括 Entry Price、Mark、Bid/Ask、数量、Greeks、价值和盈亏字段。",
        "对比：切换一个互补 Metric，例如 Notional → Total Delta，Mark IV → IV Spread。",
        "行动：记录 HOLD / CLOSE / ROLL / HEDGE / VERIFY DATA；真正下单前在券商侧复核合约、资金、报价和权限。",
    ])
    b.callout("纪律", "颜色负责排序，不负责定论。任何交易动作至少要同时满足：合约识别正确、数据状态合格、风险口径明确、流动性可接受、资金和审批可执行。", "warning")

    b.heading("4.3 Label 怎么用", 2)
    b.bullets([
        "None：默认。最适合 100–200 条仓位或完整期权链，避免数字爆炸。",
        "Held Metric：只在持仓格显示当前 Metric，适合盘中盯仓。",
        "Top 15%：只显示当前 Metric 绝对值最高的 15%，适合寻找峰值。默认在设置页隐藏。",
        "Bottom 15%：只显示当前 Metric 绝对值最低的 15%，适合找低敏感、低规模或低 IV 区域。默认在设置页隐藏。",
        "All：显示所有有值格。只建议在小范围、放大方格或做逐项复核时使用。",
    ])
    b.callout("Bottom 15% 的陷阱", "低值不等于低风险。Total Delta 很低可能是对冲抵消，也可能是数据缺失已被排除；Mark IV 很低可能是深度价内外、报价异常或模型输入问题。", "risk")

    b.heading("5. 仓位管理：输入、更新、导出", 1)
    b.heading("5.1 推荐主流程", 2)
    b.steps([
        "准备符合模板的 GLD/XAUT 持仓 Excel。Instrument、Underlying、Expiry、Strike、Call/Put、Qty、Multiplier、价格、费用和 Greeks 按列填写；Total 汇总行不当作单腿期权。",
        "在仓位管理页上传 Excel，先查看解析预览、识别出的标的和异常字段，再确认合并。不要把未识别日期、空 Strike 或错误 Call/Put 直接写入。",
        "核对导入后的持仓数、GLD/XAUT 分类、Expiry、Strike、Net Qty、Multiplier XAU 与 Entry Price。",
        "点击更新市场数据。系统刷新 Mark、Bid/Ask、IV 与相关计算；所有页面共享更新后的仓位数据。",
        "回到风险热力图按第 4 节流程检查；若数据状态为 MISSING/STALE/WARN/FAIL，先处理数据再交易。",
        "导出当前持仓 Excel。导出值应包含更新后的市场数据、价值、盈亏和计算字段，用于留档或复核。",
    ])

    b.heading("5.2 Excel 字段的业务口径", 2)
    b.bullets([
        "Qty Long、Qty Short 与 Net Qty：Net Qty 是净数量；系统风险以净数量为核心，但对运营对账应保留多空原始列。",
        "Market Value：当前 Mark 对应的仓位价值；Entry Value：买入均价对应的仓位价值；Fee 纳入 Entry Cost。",
        "Multiplier XAU：把 GLD 或 XAUT 统一换算为黄金单位。GLD 默认参数可在设置页调整，但 adjusted deliverable 必须服从合约主数据。",
        "BS Delta XAU、Gamma、Theta USD/day、Vega USD/vol：模板给出的是仓位总量时，系统需要反推或保存 Unit Greeks；缺少必要分母时不能默认为 0。",
        "Reference Date、quoteTime、positionTime、source、dataStatus：共同回答数据属于哪一天、何时更新、从哪里来、能否使用。",
    ])
    b.callout("Render 部署提醒", "当前部署若使用临时文件存储，重新部署可能清空服务端 JSON 仓位。生产环境必须迁移到持久数据库与对象存储，并保留导入文件哈希和审计记录。", "risk")

    b.heading("6. 风险计算定义与交易含义", 1)
    b.heading("6.1 价值与盈亏", 2)
    b.formula("Entry Cost = Entry Price × Net Qty × Contract Multiplier + Fee", "买入成本口径；费用正负号应按账务规范统一。")
    b.formula("Current Value = Mark Price × Net Qty × Contract Multiplier", "当前估值。负仓位自然产生负市场价值，展示和排序时需明确使用 signed 还是 absolute。")
    b.formula("UPL = Current Value − Entry Cost", "未实现盈亏；若入场与当前货币不同，还需增加 FX 换算。")
    b.formula("UPL % = UPL ÷ Entry Cost", "当 Entry Cost 为 0 或接近 0 时必须显示 MISSING/WARN，不能产生无穷或伪百分比。")

    b.heading("6.2 Notional Size 公式是否正确", 2)
    b.formula("Notional Size = Net Qty × Underlying Spot × Units of Underlying per Contract", "这个公式作为“标的名义本金”是正确的。GLD 的 Units 是实际可交割 shares；XAUT 的 Units 必须来自实时合约规格。热力图比较使用绝对值，明细保留方向符号。")
    b.p("注意：Notional 不是期权市值，也不是最大亏损，更不是 Delta-equivalent exposure。它回答“这腿对应多少标的资产”，Total Delta 才回答局部价格变动下的一阶风险。")

    b.heading("6.3 Greeks 统一为 XAU / USD", 2)
    b.formula("Spot Scale = Underlying Spot ÷ XAU/USD Spot", "把 GLD 或 XAUT 的标的单位映射为黄金单位；实际生产应优先使用合约主数据和当日 GLD ounces/share。")
    b.formula("Total Delta XAU = Unit Delta × Qty × Contract Multiplier × Spot Scale")
    b.formula("Total Gamma XAU = Unit Gamma × Qty × Contract Multiplier × Spot Scale²")
    b.formula("Total Theta USD/day = Unit Theta × Qty × Contract Multiplier")
    b.formula("Total Vega USD/vol = Unit Vega × Qty × Contract Multiplier")
    b.p("单位必须写清：某些供应商 Vega 定义为 IV 变化 1.00，某些定义为 1 vol point（0.01）；Theta 可能按自然日或交易日。跨源汇总前必须标准化，否则同名字段不能直接相加。")

    b.heading("6.4 单元格聚合与缺失规则", 2)
    b.bullets([
        "一个 Expiry × Strike 可能包含多个合约或仓位。Unit Delta、IV、IV Spread 等取代表性的最大绝对值；Total Delta、Gamma、Theta、Vega、MV、UPL、Qty、Notional 等按业务口径聚合。",
        "Qty 与 Notional 的色标只由持仓格生成，完整链中的未持仓合约不参与上下限。",
        "所选 Metric 需要的任何关键输入缺失，则该格 Metric 为 null；不染色、不打分，但保留可交易实线边框。",
        "Top/Bottom 15% 只对有有限数值的格子按绝对值分位数计算，MISSING 不会被当作 0 放入 Bottom 15%。",
    ])

    b.heading("6.5 色标算法", 2)
    b.p("默认色标是顺序风险刻度：低值绿色，中间值黄色，高值红色。系统还支持分位数、对数和零中心等缩放模式，但默认界面隐藏复杂选项。")
    b.bullets([
        "Quantile：按分位数映射，默认对极端值做 99th percentile clipping，适合厚尾仓位分布。",
        "Log：压缩数量级，适合 50 与 1,000 这类跨度明显的数据，同时保留差异。",
        "Symmetric zero-centered：以 0 为中心，适合 signed PnL 或 Delta；正负方向都重要时使用。",
        "Custom Range：为每个 Metric 单独保存上下限。跨 GLD/XAUT 做同一口径比较时，必须使用相同上下限。",
    ])

    b.heading("7. 交易与对冲场景 SOP", 1)
    b.heading("7.1 盘中 Spot 快速移动", 2)
    b.steps([
        "确认 Spot 两位小数和更新时间，不用旧 Spot 解释新风险。",
        "看 Total Delta，定位方向风险最大的持仓格；再看 Unit Delta，判断是合约敏感度还是 Qty 导致。",
        "看 IV Spread，排除“风险最大但无法经济成交”的假动作。",
        "检查最近到期日与资金约束；如果 Spot 穿越 Strike，重新评估行权、Gamma 和流动性。",
        "形成候选对冲：现货/期货/期权；比较残余风险、交易成本、滑点、资金和治理要求后执行。",
    ])

    b.heading("7.2 波动率跳升或跳降", 2)
    b.steps([
        "先看 Mark IV 热力图辨认整体抬升还是局部 skew/smile 异常。",
        "对比 Bid IV、Ask IV 与 IV Spread，判断是可成交波动率变化还是单边报价噪声。",
        "再看持仓 Vega（若设置页开放）及 Total Delta；波动率冲击常与价格方向和流动性同时变化。",
        "不要仅因高 IV 卖出期权；必须评估尾部损失、保证金、Gamma、到期路径和对冲可用性。",
    ])

    b.heading("7.3 到期日前两天", 2)
    b.steps([
        "核对 DTE、官方收盘/Spot 与 Strike 距离、合约 deliverable 和经纪商 Exercise/DNE 截止时间。",
        "估算行权资金：Strike × 实际 deliverable × contracts；adjusted contract 不能硬编码 100。",
        "比较可用 USD / Buying Power 与 Funding Coverage。可能 ITM 且覆盖率低于 100% 时应标记 FAIL。",
        "给每腿记录 Planned Action、Owner、Reviewer 与 Confirmation ID；HOLD、CLOSE、ROLL、EXERCISE、DNE 必须可追溯。",
    ])

    b.heading("8. 默认操作手册与故障排查", 1)
    b.heading("8.1 每日最小闭环", 2)
    b.steps([
        "登录后进入仓位管理，核对持仓日期与数量。",
        "点击更新市场数据，等待页面报告成功/失败与更新时间。",
        "进入风险热力图，依次看 Notional、Total Delta、Unit Delta、IV 与 IV Spread。",
        "对高风险格 Hover、点击详情并记录动作；MISSING/STALE 先修数据。",
        "完成交易或对冲后更新仓位，再刷新市场数据，确认 Total、Heatmap、Alerts 与 Scenario（若开放）同步变化。",
        "导出 Excel 并保存到受控位置，完成当日审计。",
    ])

    b.heading("8.2 常见问题", 2)
    b.table(
        ["现象", "最可能原因", "处理"],
        [
            ["方格有实线但无颜色", "可交易合约存在，但当前 Metric 缺输入", "Hover/详情看 MISSING 字段；换数据源或等待刷新，绝不手工填 0"],
            ["Top/Bottom 15% 选项看不到", "默认在设置页隐藏", "进入设置 → Heatmap filter options → Label 开启对应项"],
            ["Spot 在图外", "当前 Strike 范围不覆盖 Spot", "使用 Center Spot；或调整范围，仍保留 Above/Below 提示"],
            ["GLD 数据延迟", "Cboe delayed feed、上游时间戳或免费接口限制", "看源时间与 observed age；生产使用授权 OPRA/供应商低延迟源"],
            ["导入后数量不对", "Long/Short/Net 口径、Total 行或日期解析", "在预览阶段逐列核对，修正模板后重新导入"],
            ["部署后仓位消失", "Render 临时文件系统在部署/重启后重置", "恢复 Excel；尽快迁移 Postgres + 对象存储"],
            ["色阶跨标的无法比较", "两次自动范围不同", "固定同一 Metric 的上下限，并确认单位标准化"],
        ],
        [2100, 3220, 4040],
    )

    b.heading("8.3 安全与权限", 2)
    b.bullets([
        "不要把公网 Basic Auth、管理员密码或 API Key 写入 Excel、截图、源码或本手册；通过环境变量和密码管理器分发。",
        "风险热力图与仓位管理默认可见；其他页面按设置决定入口，且每次打开可要求管理员密码。隐藏入口不是授权控制，后端仍需 RBAC。",
        "任何公式修改都应记录修改者、时间、旧值、新值、影响范围与批准人，并能恢复默认。",
        "生产环境需启用 TLS、强密码、速率限制、会话过期、审计日志、备份与恢复演练。",
    ])

    b.heading("9. 当前已有但默认隐藏的能力", 1)
    b.p("以下能力存在于系统中或可由设置开放，但不属于默认首屏。使用前应由风险负责人决定是否对相应角色开放。")
    b.bullets([
        "额外筛选：Dataset、Venue、Broker、Account、Expiry bucket、Data Status、Color Scale、Spot source、Transpose。",
        "额外 Metric：Gamma、Theta、Vega、Market Value、UPL、DTE、Distance-to-Strike、Roll Priority 等。",
        "决策卡：Max Unit Delta、Max Total Delta、Max Theta Burn、Max Vega、Nearest Expiry、Highest Roll Priority、Largest Data Error；点击只定位高亮。",
        "Expiry / Exercise Panel：GLD DTE≤2 时检查资金覆盖、截止时间、计划动作和复核信息。",
        "Scenario：XAU Shock −20% 至 +20%、IV Shock、Day 0/1/3/7；汇总期权 PnL、Van PnL、Residual PnL 与 Coverage。",
        "Dashboard、公式管理、期权详情和其他页面：按页面入口和管理员保护配置开放。",
        "Roll Priority heuristic：由 DTE、Theta/MV、距行权价、Unit/Total Delta、spread、time value、hedge contribution、残余风险改善加权，能够展开原因与权重。",
    ])
    b.callout("不要把隐藏等同于不存在", "隐藏是信息架构选择；安全、模型治理和数据权限仍必须在后端执行。", "warning")

    b.heading("10. 当前问题、潜在提升与路线图", 1)
    b.heading("10.1 当前最需要正视的问题", 2)
    b.table(
        ["问题", "交易影响", "建议优先级与方案"],
        [
            ["Render 临时存储", "重启/部署后仓位可能丢失", "P0：Postgres 保存仓位/设置/审计；对象存储保存 Excel；定时备份"],
            ["GLD 延迟与授权", "旧报价导致错误 Delta、IV 与动作", "P0：采购 OPRA/专业供应商；记录 exchange/receive/display 三类时间"],
            ["合约主数据不足", "adjusted deliverable、乘数或到期规则错误", "P0：独立 contract master；版本化 multiplier、deliverable、exercise style、cutoff"],
            ["数据源口径不同", "Vega、Theta、IV、币种不能直接相加", "P0：数据字典、单位转换层、源字段血缘与校验"],
            ["REST 快照非流式", "快速行情时页面变化滞后", "P1：WebSocket/stream、增量计算、断线重连与心跳"],
            ["模型简化", "美式 GLD、分红、曲面、跳跃和尾部风险被低估", "P1：美式定价/曲面、利率分红、情景全重估与模型对照"],
            ["缺少执行闭环", "发现风险但未衡量滑点、冲击和订单状态", "P1：盘口深度、可成交量、pre-trade TCA、订单与成交回写"],
            ["权限偏前端", "隐藏入口不足以满足机构治理", "P0：SSO/MFA、后端 RBAC、字段级权限、审计告警"],
        ],
        [1750, 2800, 4810],
    )

    b.heading("10.2 专业化升级建议", 2)
    b.bullets([
        "PnL Explain：把日内 PnL 拆成 Delta、Gamma、Theta、Vega、Vol Surface、FX、Carry、Trade 与 Residual，帮助交易员判断“赚/亏在哪里”。",
        "Hedge Optimizer：在目标残余 Delta/Gamma/Vega、资金、保证金、成交成本和品种权限约束下，给出多个可解释候选组合，而不是让 AI 直接拍脑袋下单。",
        "Cross-market Liquidity Heatmap：统一到 USD notional、XAU delta、bps spread、1% ADV/可成交深度，比较 Bybit XAUT、GLD 与 CME；三个品种可以同屏，但必须统一量纲和时间戳。",
        "Vol Surface QA：按 expiry 构造 forward、moneyness/delta bucket、无套利检查和 stale quote mask；防止错误报价污染 Greeks。",
        "Alerts 与 Runbook：Spot 穿 Strike、0DTE、资金覆盖不足、数据中断、IV Spread 扩大、残余 Delta 超限时自动通知，并附处理手册。",
        "Position APIs：Bybit/SignalPlus 自动导入；IBKR/Futu/KGI 优先 API 或 statement parser；Excel 保留为灾备与人工复核。",
        "可观测性：接口成功率、延迟分位数、数据缺失率、重算耗时、最后成功刷新、版本号和模型哈希进入监控。",
    ])

    b.heading("10.3 建议的三阶段路线图", 2)
    b.table(
        ["阶段", "目标", "验收标准"],
        [
            ["P0：可信与可恢复", "持久化、低延迟授权源、合约主数据、RBAC、审计", "部署不丢仓；数据时间可追；缺失不静默；合约与券商一致"],
            ["P1：实时与可执行", "流式数据、API 仓位、全量情景、资金/到期、对冲候选", "行情变化秒级反映；Qty+1 后全链重算；动作能复核"],
            ["P2：优化与学习", "PnL Explain、跨市场流动性、TCA、回测、风险预算", "能量化降低了多少残余风险、成本与操作错误"],
        ],
        [1700, 3900, 3760],
    )
    b.callout("路线图使用方式", "先把 P0 的可信、持久化与治理做实，再开放 P1 的实时执行；P2 的优化价值必须用残余风险、成本和操作错误的改善来验收。", "info")

    b.heading("11. 四份面向不同人群的讲解稿", 1)
    b.heading("11.1 给最专业的师傅 / 交易同事（约 6 分钟）", 2)
    b.p("“这套系统的核心不是再做一个仓位表，而是把 GLD 与 XAUT 的期权风险投影到 Expiry × Strike 张量的一张切片。默认先看 Notional，是为了找资本和标的规模集中；切 Total Delta 看真实方向敞口；再切 Unit Delta，把合约敏感度和仓位数量拆开。这样 0.05 Delta、1,000 张与 0.80 Delta、1 张不会混淆。")
    b.p("“白框是持仓，青色细框是完整链可交易合约。颜色只表达当前 Metric 的相对大小；数据缺失就留空，不会伪装成绿色低风险。Spot 按两位小数放在最近 Strike，超范围也会明确提示。我们的读取顺序始终是：数据时间—Spot—风险簇—持仓边框—Hover—详情—流动性—动作。”")
    b.p("“做对冲时，不把热力图当 optimizer。先用 Total Delta/Gamma/Vega 找残余风险，再加入 spread、深度、资金、保证金和到期规则。Roll Priority 只是透明 heuristic：DTE、Theta/MV、moneyness、Delta、spread、time value、hedge contribution 与 residual improvement 都能解释。最终动作仍然要有候选比较、双人复核和成交回写。”")
    b.p("“当前最大的技术风险是数据与持久化：GLD 延迟源不能冒充实时，Render 临时存储不能冒充账本，adjusted contract 不能用 100 shares 硬编码。下一步应先完成 OPRA/专业源、contract master、Postgres、审计和 RBAC，再做流式曲面与 hedge optimizer。”")
    b.callout("给专业同事的结束句", "这张图的价值不是让我们看得更漂亮，而是让风险发现、数据质疑、对冲选择和事后复盘使用同一套坐标与口径。", "positive")

    b.heading("11.2 给黄金交易所 / 公司大老板（约 3 分钟）", 2)
    b.p("“我们把两类黄金相关资产——传统市场的 GLD 与数字资产市场的 XAUT——放进一套统一风险语言。管理层不需要读几百行期权，只要看颜色、边框和 Spot，就能知道风险集中在哪个到期日与行权价、哪些是真实持仓、哪些数据不够新。”")
    b.p("“这套工具最直接的经营价值有四个：第一，更快识别大风险；第二，减少把单张 Delta 当成仓位风险的判断错误；第三，在到期、行权和资金覆盖上提前暴露 FAIL；第四，把数据来源、时间、公式和操作人纳入审计，让风险流程可复核。”")
    b.p("“它目前是风险决策支持平台，不是交易所账本，也不会自动下单。要达到机构生产级，投资顺序不是先做更多图，而是低延迟授权数据、合约主数据、持久数据库、权限审计和灾备。完成这些底座后，才能可靠扩展跨市场流动性比较、自动对冲建议和 PnL 解释。”")
    b.callout("管理层指标", "建议按月跟踪：数据可用率、P95 延迟、MISSING 比例、风险超限发现时间、到期失败数、对冲成本、部署恢复时间与审计覆盖率。", "info")

    b.heading("11.3 给公司其他部门同事（约 5 分钟）", 2)
    b.p("“可以把热力图理解为一张期权风险地图。横向是不同到期日，纵向是不同行权价。每个小格子是一类合约；绿色到红色代表我们当前选择的指标从低到高，白色边框表示公司真的持有，青色边框表示市场上可交易但没有持有。”")
    b.p("“默认指标 Notional Size 讲的是这笔期权对应多少标的资产，不等于花了多少钱，也不等于最大亏损。切到 Total Delta 才更接近小幅金价变化时的方向风险。切到 IV 和 IV Spread，能看到市场对未来波动的定价和交易难度。”")
    b.p("“运营同事最重要的是数据链：仓位 Excel 导入是否正确、市场数据更新时间、缺失和警告是否处理、导出是否留档。技术同事最重要的是接口、持久化、权限和审计。财务同事最重要的是币种、成本、费用、市场价值、盈亏和资金覆盖口径。”")
    b.callout("跨部门共同语言", "一个结论至少带上：标的、到期日、行权价、Call/Put、当前 Metric、时间戳、数据状态和负责人。", "positive")

    b.heading("11.4 给小白（约 8 分钟）", 2)
    b.p("“期权可以理解为一张在未来某天、按某个价格买或卖标的的权利。Expiry 是这张权利什么时候到期，Strike 是约定价格，Call 是看涨方向的权利，Put 是看跌方向的权利。”")
    b.p("“热力图把很多张期权排成格子：横着找日期，竖着找价格。白边格是我们手里有的，青边格是市场有但我们没持有的。颜色不是涨跌预测，而是当前指标的大小。默认 Notional Size 只是规模；要看金价变化的敏感度，就选 Delta。”")
    b.p("“Unit Delta 是一张期权的敏感程度，Total Delta 是把张数也算进去。比如一张 0.8 Delta 的期权看起来很敏感，但只有 1 张；另一种只有 0.05 Delta，却有 1,000 张，后者的总风险可能远大于前者。所以一定要看清 Unit 还是 Total。”")
    b.p("“Spot 是现在金价或标的价格。网站把它显示到小数点后两位，并告诉你离哪个 Strike 最近。鼠标放到格子上可以看关键数据，点击可以看完整数据。看到红色不能马上交易，要先看数据是不是 LIVE、有没有 MISSING、买卖价差大不大。”")
    b.p("“最安全的使用顺序是：先更新数据，再看 Spot，再看白框红格，再 Hover，看不懂就点详情；最后把判断交给有权限的交易员复核。热力图帮助你找到该问的问题，但不会替你承担交易决策。”")
    b.callout("给小白的口诀", "先认坐标，再认颜色；先看数据，再看风险；先问口径，再做动作。", "info")

    b.heading("12. 第一性原理总结", 1)
    b.heading("12.1 从风险本质出发", 2)
    b.p("期权仓位的本质不是“有多少张”，而是在价格、波动率、时间和流动性变化下，未来现金流如何改变。Greeks 是局部导数，是复杂损益曲面在当下附近的近似；因此 Unit 与 Total、局部与情景、模型与市场都必须分开。")
    b.formula("Risk ≈ Delta·ΔS + ½·Gamma·(ΔS)² + Vega·ΔIV + Theta·Δt + Residual", "这只是局部近似。价格跳跃、波动率曲面移动、流动性枯竭和行权规则会进入 Residual。")

    b.heading("12.2 从热力图本质出发", 2)
    b.p("原始风险是 Underlying × Venue × Account × Call/Put × Expiry × Strike × Metric × Time 的高维稀疏张量。热力图选择 Expiry × Strike 作为二维切片，用筛选器确定其余维度，用颜色压缩数值，用边框编码持仓状态，用角标编码合约和数据状态。它的目标不是展示更多，而是用最少视觉变量保留最关键决策信息。")

    b.heading("12.3 从对冲本质出发", 2)
    b.p("对冲不是把所有风险降到 0，而是在收益目标、风险限额、交易成本、流动性、资金、保证金、税务和治理约束下，选择可承受的残余风险。最优动作必须明确优化目标和约束，不能只按一个红色格子下单。")
    b.formula("Choose hedge h to minimize: Residual Risk(h) + Transaction Cost(h) + Funding Cost(h)", "subject to liquidity、position limits、margin、exercise rules、approvals 与可解释性。")

    b.heading("12.4 从数据本质出发", 2)
    b.p("任何风险数字都是 Data × Convention × Model × Time 的产物。数据源、时间戳、合约定义和单位约定先决定数字是否有意义，模型精度才决定数字有多准确。因此 MISSING 不能等于 0，STALE 不能等于 LIVE，adjusted contract 不能等于标准合约。")

    b.heading("12.5 从系统治理本质出发", 2)
    b.p("机构级风险系统必须让每个结论可重复、可归因、可恢复、可审计。公式要版本化，市场数据要有血缘，仓位要能对账，权限要在后端执行，部署要能恢复，动作要有 Owner/Reviewer/Confirmation。否则再漂亮的热力图也只是不可控的单机工具。")

    b.heading("12.6 分别面向四类人的一句话", 2)
    b.bullets([
        "对专业交易员：先把高维风险压缩并排序，再用流动性、情景和约束决定残余风险。",
        "对管理层：把风险发现时间、数据可信度、资金失败与审计覆盖率变成可管理指标。",
        "对其他部门：用统一字段、时间戳和责任人把交易、技术、运营、财务连接起来。",
        "对小白：热力图是找问题的地图，不是预测涨跌的答案。",
    ])
    b.callout("第一性原理总括", "价格风险决定看什么；数据血缘决定能不能信；流动性、资金与治理约束决定能不能做；残余风险与成本的改善决定做得值不值。", "positive")

    b.page_break()
    b.heading("附录 A：默认状态速查", 1)
    b.table(
        ["项目", "当前默认"],
        [
            ["默认页面入口", "风险热力图、仓位管理、设置"],
            ["Underlying", "GLD"],
            ["Call / Put", "Call"],
            ["Metric", "Notional Size"],
            ["Label", "None；Top 15% 与 Bottom 15% 默认隐藏"],
            ["Hover", "Risk"],
            ["颜色", "低绿—中黄—高红；缺失留空"],
            ["持仓边框", "白色较粗实线"],
            ["可交易未持仓边框", "青色细实线，四边完整"],
            ["Spot", "最近 Strike 标记；数值两位小数；越界仍提示"],
            ["默认可见控件", "Underlying、C/P、Metric、Label、Hover、Range、Strike 排序、方格尺寸、Fit All、Full Screen"],
            ["默认隐藏区", "情景分析、Show Cards、Largest Data Error 与其他高级筛选"],
        ],
        [2500, 6860],
    )

    b.heading("附录 B：术语速查", 1)
    b.table(
        ["术语", "简明解释"],
        [
            ["Mark Price", "用于估值的代表价格，不保证等于可成交价格"],
            ["IV", "由期权价格反推出的隐含波动率"],
            ["Delta", "标的价格小幅变化时，期权价值的一阶敏感度"],
            ["Gamma", "Delta 随标的价格变化的速度"],
            ["Theta", "其他条件不变时，时间流逝带来的价值变化"],
            ["Vega", "隐含波动率变化时，期权价值的敏感度"],
            ["DTE", "距到期剩余天数"],
            ["MV", "Market Value，当前市场价值"],
            ["UPL", "Unrealized PnL，未实现盈亏"],
            ["Notional", "对应标的资产的名义本金，不是期权市值"],
            ["STALE", "报价超过允许时效；当前系统阈值按配置/实现判断"],
            ["MISSING", "必要输入缺失，不能用于计算该 Metric"],
        ],
        [2100, 7260],
    )

    b.heading("附录 C：交易前检查清单", 1)
    b.bullets([
        "仓位日期、数量、合约、乘数、币种与券商一致。",
        "市场数据 Source、As-of、时区、状态可接受。",
        "Spot、Mark、Bid/Ask、IV 与 Greeks 没有静默 0。",
        "当前看的究竟是 Unit 还是 Total；究竟是 signed 还是 absolute。",
        "颜色上下限是否自动；跨标的比较是否固定同一 Metric 的范围。",
        "到期、行权、DNE、资金覆盖、买入能力和截止时间已复核。",
        "流动性、spread、深度、滑点、手续费、保证金已计入。",
        "动作、Owner、Reviewer、Confirmation 与导出留档齐全。",
    ])
    b.callout("最终原则", "先保证数据和口径正确，再追求模型精度；先明确残余风险目标，再讨论对冲工具；先形成可复核流程，再增加自动化。", "positive")

    return doc


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    document = build_manual()
    document.save(OUTPUT_FILE)
    print(OUTPUT_FILE)


if __name__ == "__main__":
    main()
